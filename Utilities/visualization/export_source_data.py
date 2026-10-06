#!/usr/bin/env python3
"""
Export the Source Data file requested by Nature Communications.

One sheet per table or graph, each starting with a title and a short description:
    Table 1a (WSI) and Table 1b (STP): per-tile counts and scores for every
        method, plus the pooled row reported in the table. The supplementary
        bar plots show the same data.
    Supplementary persistence and min-size sweeps, for each dataset.

Writes results/tables/Source_Data.xlsx. If openpyxl is not installed, writes
one CSV per sheet into results/tables/Source_Data/ and zips the folder instead.
"""

import os
import shutil
import pandas as pd

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
BASE_DIR = os.path.dirname(os.path.dirname(SCRIPT_DIR))
OUTPUTS_DIR = os.path.join(BASE_DIR, 'outputs')
TABLES_DIR = os.path.join(BASE_DIR, 'results', 'tables')

METHODS = [
    ('bwskel', 'bwskel_evaluation'),
    ('Diff. Skel', 'diffskel_evaluation'),
    ('neuTube', 'neutube_evaluation'),
    ('VESS', 'vess_evaluation'),
    ('PHDF', 'phd_evaluation'),
    ('DM2D', 'dm2d_evaluation'),
]

COLUMNS = ['Method', 'Tile', 'TP', 'FP', 'FN', 'Precision', 'Recall', 'F1', 'IoU']


def scores(tp, fp, fn):
    precision = tp / (tp + fp) if tp + fp else 0.0
    recall = tp / (tp + fn) if tp + fn else 0.0
    f1 = 2 * precision * recall / (precision + recall) if precision + recall else 0.0
    iou = tp / (tp + fp + fn) if tp + fp + fn else 0.0
    return precision, recall, f1, iou


def table1_sheet(dataset):
    rows = []
    for method, folder in METHODS:
        path = os.path.join(OUTPUTS_DIR, dataset, folder, 'per_image_metrics.csv')
        if not os.path.exists(path):
            print(f"Warning: {path} not found - {method} left out")
            continue
        df = pd.read_csv(path).sort_values('filename')
        for r in df.itertuples():
            rows.append([method, os.path.splitext(r.filename)[0], r.TP, r.FP, r.FN,
                         r.precision, r.recall, r.f_score, r.iou])
        tp, fp, fn = df['TP'].sum(), df['FP'].sum(), df['FN'].sum()
        rows.append([method, f'All tiles pooled (n = {len(df)})', tp, fp, fn, *scores(tp, fp, fn)])
        per_tile = df[['precision', 'recall', 'f_score', 'iou']]
        rows.append([method, f'Mean across tiles (n = {len(df)})', None, None, None, *per_tile.mean()])
        rows.append([method, f'SD across tiles (n = {len(df)})', None, None, None, *per_tile.std(ddof=1)])
    return pd.DataFrame(rows, columns=COLUMNS)


def sweep_sheet(path, parameter):
    if not os.path.exists(path):
        print(f"Warning: {path} not found")
        return None
    df = pd.read_csv(path)
    return df.rename(columns={df.columns[0]: parameter, 'precision': 'Precision', 'recall': 'Recall',
                              'f_score': 'F1', 'iou': 'IoU'})


def build_sheets():
    pixel_rule = ('TP = detected skeleton pixels within the match distance of a ground-truth pixel; '
                  'FP = detected pixels with no ground-truth pixel within that distance; '
                  'FN = ground-truth pixels with no detected pixel within that distance. '
                  'The pooled row sums TP, FP and FN over all tiles before computing the scores (the value in Table 1); '
                  'mean and SD (sample standard deviation) are taken over the per-tile scores.')
    sheets = []
    for dataset, label in (('pmd', 'Table 1a (WSI)'), ('stp', 'Table 1b (STP)')):
        sheets.append((label, f'{label}: skeletonization accuracy per tile and pooled, for each method. '
                              f'The supplementary bar plot for this dataset shows the same data. {pixel_rule}',
                       table1_sheet(dataset)))
    for dataset, name in (('pmd', 'WSI'), ('stp', 'STP')):
        sweeps = [
            (f'Supp persistence {name}', 'Persistence threshold',
             f'DM2D accuracy on the {name} tiles as the persistence threshold varies (all tiles pooled).',
             os.path.join(OUTPUTS_DIR, dataset, 'dm2d_persistence_sweep', 'persistence_sweep.csv')),
            (f'Supp min-size {name}', 'Minimum component size (pixels)',
             f'DM2D accuracy on the {name} tiles as the minimum component size varies (all tiles pooled).',
             os.path.join(OUTPUTS_DIR, dataset, 'dm2d_min_size_sweep', 'min_size_sweep.csv')),
        ]
        for sheet, parameter, description, path in sweeps:
            df = sweep_sheet(path, parameter)
            if df is not None:
                sheets.append((sheet, description, df))
    return sheets


def write_xlsx(sheets, path):
    from openpyxl.styles import Font

    with pd.ExcelWriter(path, engine='openpyxl') as writer:
        for name, description, df in sheets:
            df.to_excel(writer, sheet_name=name, index=False, startrow=3)
            ws = writer.sheets[name]
            ws['A1'] = name
            ws['A1'].font = Font(size=12)
            ws['A2'] = description
            for column_cells in ws.columns:
                width = max(len(str(c.value)) for c in column_cells[3:] if c.value is not None)
                ws.column_dimensions[column_cells[0].column_letter].width = min(max(width + 2, 8), 40)
    print(f"Saved: {path}")


def write_csv_zip(sheets, folder):
    os.makedirs(folder, exist_ok=True)
    for name, description, df in sheets:
        csv_path = os.path.join(folder, name.replace(' ', '_').replace('(', '').replace(')', '') + '.csv')
        with open(csv_path, 'w') as f:
            f.write(f'# {name}\n# {description}\n')
            df.to_csv(f, index=False)
    shutil.make_archive(folder, 'zip', folder)
    print(f"Saved: {folder}.zip (openpyxl not installed, wrote CSV files instead)")


def main():
    os.makedirs(TABLES_DIR, exist_ok=True)
    sheets = build_sheets()
    if not sheets:
        print("No results found - nothing to export")
        return
    try:
        import openpyxl  # noqa: F401
        write_xlsx(sheets, os.path.join(TABLES_DIR, 'Source_Data.xlsx'))
    except ImportError:
        write_csv_zip(sheets, os.path.join(TABLES_DIR, 'Source_Data'))


if __name__ == '__main__':
    main()
